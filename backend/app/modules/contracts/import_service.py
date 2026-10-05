"""Validação e substituição transacional de contratos a partir de planilhas."""

from __future__ import annotations

import csv
import hashlib
import io
import unicodedata
import zipfile
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from itertools import islice
from pathlib import PurePath
from typing import Any

from fastapi import UploadFile
from openpyxl import load_workbook
from openpyxl.utils.exceptions import InvalidFileException
from sqlalchemy import delete, func, select, text
from sqlalchemy.ext.asyncio import AsyncSession

from app.core.auth import CurrentUser
from app.core.errors import ConflictError, DomainError, NotFoundError
from app.models.contracts import (
    Contract,
    ContractCriticality,
    OverdueContractTracking,
    OverdueDailySnapshot,
    Sector,
    Supplier,
)
from app.modules.contracts.overdue_history import _RECONCILE_LOCK_KEY
from app.modules.contracts.rules import contracts_today
from app.shared.audit import record_audit

MAX_FILE_BYTES = 5 * 1024 * 1024
MAX_UNCOMPRESSED_BYTES = 50 * 1024 * 1024
MAX_ROWS = 10000
HEADERS = (
    "Contrato",
    "Fornecedor",
    "Prestação",
    "Valor Serviço",
    "Valor Material Próprio",
    "Valor Material Terceiros",
    "Valor Total",
    "Início Vigência",
    "Fim Vigência",
    "Unidade",
    "Situação",
    "Criticidade",
    "Renovação Automática",
)
_REQUIRED = {"contrato", "fornecedor"}
_MONEY = ("valor servico", "valor material proprio", "valor material terceiros", "valor total")
_DATES = ("inicio vigencia", "fim vigencia")
_MAX_MONEY = Decimal("99999999999999.99")


def _clean(value: Any) -> str:
    return " ".join(str(value).split()) if value is not None else ""


def _key(value: Any) -> str:
    normalized = unicodedata.normalize("NFKD", _clean(value).casefold())
    return "".join(char for char in normalized if not unicodedata.combining(char))


def _issue(line: int, column: str, reason: str) -> dict[str, Any]:
    return {"line": line, "column": column, "reason": reason}


def _money(value: Any) -> Decimal:
    if value is None or _clean(value) == "":
        return Decimal("0")
    if isinstance(value, bool):
        raise ValueError("Valor financeiro inválido")
    raw = _clean(value)
    if isinstance(value, str):
        raw = raw.replace("R$", "").replace(" ", "")
        if "," in raw:
            raw = raw.replace(".", "").replace(",", ".")
        elif raw.count(".") > 1:
            raise ValueError("Valor financeiro inválido")
    try:
        amount = Decimal(raw)
    except InvalidOperation as exc:
        raise ValueError("Valor financeiro inválido") from exc
    if not amount.is_finite():
        raise ValueError("Valor financeiro inválido")
    if amount < 0 or amount > _MAX_MONEY or amount != amount.quantize(Decimal("0.01")):
        raise ValueError("Valor financeiro inválido (use até duas casas decimais, sem valor negativo)")
    return amount.quantize(Decimal("0.01"))


def _date(value: Any) -> date | None:
    if value is None or _clean(value) == "":
        return None
    if isinstance(value, datetime):
        return value.date()
    if isinstance(value, date):
        return value
    raw = _clean(value)
    for format_string in ("%d/%m/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(raw, format_string).date()
        except ValueError:
            pass
    raise ValueError("Data inválida; use dd/mm/aaaa")


def _bool(value: Any) -> bool:
    if value is None or _clean(value) == "":
        return False
    if isinstance(value, bool):
        return value
    normalized = _key(value)
    if normalized in {"sim", "s", "true", "1"}:
        return True
    if normalized in {"nao", "n", "false", "0"}:
        return False
    raise ValueError("Use Sim ou Não")


@dataclass(slots=True)
class ImportRow:
    number: str
    supplier: str
    description: str
    service_value: Decimal
    own_value: Decimal
    third_value: Decimal
    total_value: Decimal
    start_date: date | None
    end_date: date | None
    unit: str
    finalized: bool
    criticality: ContractCriticality | None
    auto_renewal: bool


@dataclass(slots=True)
class ParsedFile:
    filename: str
    sha256: str
    source: str
    found: int = 0
    rows: list[ImportRow] = field(default_factory=list)
    issues: list[dict[str, Any]] = field(default_factory=list)
    duplicates: int = 0


def _parse_row(values: dict[str, Any]) -> tuple[ImportRow | None, list[dict[str, Any]]]:
    line = int(values["_line"])
    issues: list[dict[str, Any]] = []
    number = _clean(values.get("contrato"))
    supplier = _clean(values.get("fornecedor"))
    if not number:
        issues.append(_issue(line, "Contrato", "Campo obrigatório"))
    if not supplier:
        issues.append(_issue(line, "Fornecedor", "Campo obrigatório"))
    if len(number) > 80:
        issues.append(_issue(line, "Contrato", "Máximo de 80 caracteres"))
    if len(supplier) > 240:
        issues.append(_issue(line, "Fornecedor", "Máximo de 240 caracteres"))
    unit = _clean(values.get("unidade"))
    if len(unit) > 120:
        issues.append(_issue(line, "Unidade", "Máximo de 120 caracteres"))

    converted: dict[str, Any] = {}
    for name in _MONEY:
        try:
            converted[name] = _money(values.get(name))
        except ValueError as exc:
            issues.append(_issue(line, name, str(exc)))
    for name in _DATES:
        try:
            converted[name] = _date(values.get(name))
        except ValueError as exc:
            issues.append(_issue(line, name, str(exc)))
    situation = _key(values.get("situacao"))
    if situation not in {"", "vigente", "vencido", "finalizado"}:
        issues.append(_issue(line, "Situação", "Use Vigente, Vencido ou Finalizado"))
    criticality = _key(values.get("criticidade"))
    if criticality not in {"", "baixa", "media", "alta"}:
        issues.append(_issue(line, "Criticidade", "Use Baixa, Média ou Alta"))
    try:
        renewal = _bool(values.get("renovacao automatica"))
    except ValueError as exc:
        issues.append(_issue(line, "Renovação Automática", str(exc)))
        renewal = False
    start = converted.get("inicio vigencia")
    end = converted.get("fim vigencia")
    if start and end and start > end:
        issues.append(_issue(line, "Fim Vigência", "Fim anterior ao início da vigência"))
    total = converted.get("valor total")
    if values.get("valor total") is None or _clean(values.get("valor total")) == "":
        total = sum((converted.get(name, Decimal("0")) for name in _MONEY[:3]), Decimal("0"))
        if total > _MAX_MONEY:
            issues.append(_issue(line, "Valor Total", "Valor total excede o limite"))
    if issues:
        return None, issues
    assert isinstance(total, Decimal)
    return ImportRow(
        number=number,
        supplier=supplier,
        description=_clean(values.get("prestacao")),
        service_value=converted["valor servico"],
        own_value=converted["valor material proprio"],
        third_value=converted["valor material terceiros"],
        total_value=total,
        start_date=start,
        end_date=end,
        unit=unit,
        finalized=situation == "finalizado",
        criticality=ContractCriticality(criticality.upper()) if criticality else None,
        auto_renewal=renewal,
    ), []


def _table(content: bytes, extension: str) -> list[list[Any]]:
    if extension == ".csv":
        try:
            decoded = content.decode("utf-8-sig")
        except UnicodeDecodeError as exc:
            raise DomainError("CSV deve estar em UTF-8") from exc
        try:
            delimiter = csv.Sniffer().sniff(decoded[:8192], delimiters=",;").delimiter
        except csv.Error:
            # Cabeçalho com uma única coluna ainda deve chegar ao preview para
            # apontar as colunas obrigatórias ausentes.
            first_line = decoded.splitlines()[0] if decoded.splitlines() else ""
            delimiter = ";" if first_line.count(";") >= first_line.count(",") else ","
        try:
            return list(csv.reader(io.StringIO(decoded), delimiter=delimiter))
        except csv.Error as exc:
            raise DomainError("CSV inválido") from exc
    try:
        with zipfile.ZipFile(io.BytesIO(content)) as archive:
            if sum(info.file_size for info in archive.infolist()) > MAX_UNCOMPRESSED_BYTES:
                raise DomainError("XLSX descompactado excede 50 MB")
        workbook = load_workbook(io.BytesIO(content), read_only=True, data_only=False, keep_links=False)
        try:
            sheet = workbook.active
            if sheet is None:
                return []
            return [list(row) for row in islice(sheet.iter_rows(values_only=True), MAX_ROWS + 2)]
        finally:
            workbook.close()
    except (OSError, ValueError, zipfile.BadZipFile, KeyError, InvalidFileException) as exc:
        raise DomainError("XLSX inválido") from exc


async def parse_upload(upload: UploadFile) -> ParsedFile:
    filename = PurePath((upload.filename or "").replace("\\", "/")).name
    extension = "." + filename.rsplit(".", 1)[-1].lower() if "." in filename else ""
    if extension not in {".xlsx", ".csv"}:
        raise DomainError("Envie um arquivo .xlsx ou .csv")
    allowed_types = {
        ".xlsx": {
            "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
            "application/octet-stream",
            "",
        },
        ".csv": {"text/csv", "text/plain", "application/vnd.ms-excel", "application/octet-stream", ""},
    }
    if upload.content_type not in allowed_types[extension]:
        raise DomainError("Tipo de arquivo não permitido")
    content = await upload.read(MAX_FILE_BYTES + 1)
    if len(content) > MAX_FILE_BYTES:
        raise DomainError("Arquivo excede 5 MB")
    parsed = ParsedFile(
        filename=filename, sha256=hashlib.sha256(content).hexdigest(), source=f"import_{extension[1:]}"
    )
    table = _table(content, extension)
    if not table:
        parsed.issues.append(_issue(1, "Arquivo", "Planilha vazia"))
        return parsed
    headers = [_key(header) for header in table[0]]
    if len(headers) != len(set(headers)):
        parsed.issues.append(_issue(1, "Cabeçalho", "Colunas duplicadas"))
    for required in sorted(_REQUIRED - set(headers)):
        parsed.issues.append(_issue(1, required, "Coluna obrigatória ausente"))
    if parsed.issues:
        parsed.found = sum(any(_clean(value) for value in cells) for cells in table[1:])
        return parsed
    seen: set[str] = set()
    for line, cells in enumerate(table[1:], start=2):
        if line > MAX_ROWS + 1:
            parsed.issues.append(_issue(line, "Arquivo", "Máximo de 10.000 linhas"))
            break
        if not any(_clean(value) for value in cells):
            continue
        parsed.found += 1
        if any(_clean(value) for value in cells[len(headers) :]):
            parsed.issues.append(_issue(line, "Arquivo", "Linha contém mais colunas que o cabeçalho"))
            continue
        values = dict(zip(headers, cells, strict=False))
        values["_line"] = line
        formula_columns = [header for header, value in values.items() if _clean(value).startswith("=")]
        if formula_columns:
            parsed.issues.extend(
                _issue(line, column, "Fórmulas não são permitidas") for column in formula_columns
            )
            continue
        row, row_issues = _parse_row(values)
        number_key = _key(values.get("contrato"))
        if number_key and number_key in seen:
            parsed.duplicates += 1
            row_issues.append(_issue(line, "Contrato", "Número duplicado na planilha"))
            row = None
        seen.add(number_key)
        parsed.issues.extend(row_issues)
        if row is not None:
            parsed.rows.append(row)
    if not parsed.found:
        parsed.issues.append(_issue(2, "Arquivo", "Nenhum contrato encontrado"))
    return parsed


def _fingerprint(contracts: list[tuple[str, datetime]]) -> str:
    payload = "\n".join(f"{contract_id}:{updated.isoformat()}" for contract_id, updated in sorted(contracts))
    return hashlib.sha256(payload.encode()).hexdigest()


async def _state(session: AsyncSession, sector_id: str) -> tuple[int, str, int, bool]:
    sector = await session.get(Sector, sector_id)
    if sector is None or not sector.active:
        raise NotFoundError("Setor não encontrado")
    contracts = (
        await session.execute(select(Contract.id, Contract.updatedAt).where(Contract.sectorId == sector_id))
    ).all()
    other_count = (
        await session.execute(
            select(func.count()).select_from(Contract).where(Contract.sectorId != sector_id)
        )
    ).scalar_one()
    has_history = (await session.execute(select(OverdueDailySnapshot.id).limit(1))).first() is not None
    has_tracking = (await session.execute(select(OverdueContractTracking.id).limit(1))).first() is not None
    return (
        len(contracts),
        _fingerprint([(row.id, row.updatedAt) for row in contracts]),
        other_count,
        has_history or has_tracking,
    )


async def preview(session: AsyncSession, sector_id: str, parsed: ParsedFile) -> dict[str, Any]:
    current_count, fingerprint, other_count, has_history = await _state(session, sector_id)
    issues = list(parsed.issues)
    if other_count and has_history:
        issues.append(
            _issue(
                0,
                "Histórico",
                "Há histórico corporativo e contratos em outros setores. "
                "A substituição parcial está bloqueada para preservar o histórico global.",
            )
        )
    existing = (await session.execute(select(Supplier.name))).scalars().all()
    existing_keys = {_key(name) for name in existing}
    new_suppliers = {_key(row.supplier) for row in parsed.rows if _key(row.supplier) not in existing_keys}
    today = contracts_today()
    return {
        "filename": parsed.filename,
        "fileSha256": parsed.sha256,
        "sectorId": sector_id,
        "sectorFingerprint": fingerprint,
        "currentContracts": current_count,
        "rowsFound": parsed.found,
        "validRows": len(parsed.rows),
        "invalidRows": parsed.found - len(parsed.rows),
        "contractsToImport": len(parsed.rows),
        "newSuppliers": len(new_suppliers),
        "units": sorted({row.unit for row in parsed.rows if row.unit}),
        "overdueContracts": sum(
            not row.finalized and row.end_date is not None and row.end_date < today for row in parsed.rows
        ),
        "finalizedContracts": sum(row.finalized for row in parsed.rows),
        "duplicates": parsed.duplicates,
        "issues": issues,
        "canConfirm": not issues,
        "historyAction": (
            "blocked"
            if other_count and has_history
            else "create_baseline"
            if other_count
            else "reset_baseline"
        ),
    }


async def confirm(
    session: AsyncSession,
    actor: CurrentUser,
    sector_id: str,
    parsed: ParsedFile,
    expected_sha256: str,
    expected_fingerprint: str,
) -> dict[str, Any]:
    if parsed.sha256 != expected_sha256:
        raise ConflictError("Arquivo mudou após o preview. Valide novamente")
    if parsed.issues:
        raise DomainError("Planilha contém erros. Corrija e valide novamente")
    # O mesmo lock da reconciliação protege snapshots. O lock por setor torna
    # confirmações simultâneas seriadas, inclusive quando o setor está vazio.
    await session.execute(text("SELECT pg_advisory_xact_lock(:key)"), {"key": _RECONCILE_LOCK_KEY})
    sector = (
        await session.execute(select(Sector).where(Sector.id == sector_id).with_for_update())
    ).scalar_one_or_none()
    if sector is None or not sector.active:
        raise NotFoundError("Setor não encontrado")
    current_count, fingerprint, other_count, has_history = await _state(session, sector_id)
    if fingerprint != expected_fingerprint:
        raise ConflictError("Contratos do setor mudaram após o preview. Valide novamente")
    if other_count and has_history:
        raise ConflictError("Histórico corporativo de outros setores impede a substituição parcial")
    # Os contratos do setor são bloqueados antes do DELETE; uma edição concorrente
    # espera a transação terminar e não produz um estado parcialmente importado.
    await session.execute(select(Contract.id).where(Contract.sectorId == sector_id).with_for_update())
    await session.execute(delete(Contract).where(Contract.sectorId == sector_id))
    if not other_count:
        await session.execute(delete(OverdueDailySnapshot))
    suppliers = (
        (await session.execute(select(Supplier).order_by(Supplier.createdAt, Supplier.id))).scalars().all()
    )
    by_name = {_key(supplier.name): supplier for supplier in suppliers}
    created_suppliers = 0
    for row in parsed.rows:
        key = _key(row.supplier)
        supplier = by_name.get(key)
        if supplier is None:
            supplier = Supplier(name=row.supplier, active=True)
            session.add(supplier)
            await session.flush()
            by_name[key] = supplier
            created_suppliers += 1
        session.add(
            Contract(
                sectorId=sector_id,
                supplierId=supplier.id,
                contractNumber=row.number,
                serviceDescription=row.description,
                serviceValue=row.service_value,
                ownMaterialValue=row.own_value,
                thirdPartyMaterialValue=row.third_value,
                totalValue=row.total_value,
                startDate=row.start_date,
                endDate=row.end_date,
                unit=row.unit,
                finalized=row.finalized,
                criticality=row.criticality,
                autoRenewal=row.auto_renewal,
                source=parsed.source,
                notify=False,
            )
        )
    await session.flush()
    # Uma nova baseline começa no mesmo commit. Para uma base parcial só chegamos
    # aqui quando não existia histórico anterior, logo incluímos todos os setores.
    overdue = (await session.execute(select(Contract.id, Contract.endDate, Contract.finalized))).all()
    today = contracts_today()
    overdue_ids = [
        item.id
        for item in overdue
        if not item.finalized and item.endDate is not None and item.endDate < today
    ]
    session.add_all(
        OverdueContractTracking(contractId=contract_id, firstSeenOverdueAt=today)
        for contract_id in overdue_ids
    )
    session.add(OverdueDailySnapshot(date=today, remaining=len(overdue_ids), resolvedCumulative=0))
    await record_audit(
        session,
        user_id=actor.id,
        action="contract.import.replace",
        entity="Sector",
        entity_id=sector_id,
        metadata={
            "sectorId": sector_id,
            "filename": parsed.filename,
            "sha256": parsed.sha256,
            "previousCount": current_count,
            "importedCount": len(parsed.rows),
            "createdSuppliers": created_suppliers,
            "mode": "replace",
        },
    )
    await session.commit()
    return {
        "previousCount": current_count,
        "importedCount": len(parsed.rows),
        "createdSuppliers": created_suppliers,
    }
